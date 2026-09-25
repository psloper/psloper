#!/usr/bin/env python3
"""Merge the WINDEL UK Wind Energy Locations extract into data/uk-wind-farms.json.

WHY A MERGE AND NOT A REPLACEMENT

Both sources are the same underlying database, the DESNZ Renewable Energy
Planning Database, reached two different ways:

  - data/uk-wind-farms.json comes from a GitHub mirror of a later REPD
    snapshot. It is more current but undocumented.
  - data/windel/REPD_202407_WIND.xlsx is the July 2024 release, compiled and
    documented by Datadaptive under the Open Government Licence. It is older
    but it carries three fields the mirror drops entirely: the number of
    turbines, their height, and their individual capacity.

MEASURED, on the 2,371 REPD references both carry: the positions agree to a
median of 1.7 m, which is transform rounding, so neither is more accurate than
the other. But each source holds live projects the other does not: 115 live
only in the mirror (mostly applications submitted after July 2024, as you would
expect from the later snapshot) and 172 live only in WINDEL. A union is
strictly better than either alone, so that is what this does.

Rules:
  - position and status come from the mirror where it has the record, because
    it is the later snapshot; from WINDEL otherwise.
  - turbine count, tip height and per-turbine capacity come from WINDEL, which
    is the only source for them.
  - where the two sources' positions disagree by more than 100 m, that
    disagreement is recorded. It is evidence about the reliability of the
    point, and hiding it would throw that evidence away.

WHAT THIS DOES NOT FIX: positional accuracy. Against the two sites where this
tool has ground truth, WINDEL gives 1,140 m at Kelmarsh and 1,121 m at
Penmanshiel, which are the same errors the mirror gives to the metre. Both are
planning-application grid references, and neither is a turbine position.

Run from the windfarm-radar directory:  python3 calibration/merge_windel.py
"""
import json, math, os, sys

XLSX = 'data/windel/REPD_202407_WIND.xlsx'
OUT = 'data/uk-wind-farms.json'
LIVE = ("Operational", "Under Construction", "Awaiting Construction", "Application Submitted")


def great_circle_m(a_lat, a_lon, b_lat, b_lon):
    r = 6371008.8
    p1, p2 = math.radians(a_lat), math.radians(b_lat)
    h = (math.sin((p2 - p1) / 2) ** 2
         + math.cos(p1) * math.cos(p2) * math.sin(math.radians(b_lon - a_lon) / 2) ** 2)
    return 2 * r * math.asin(math.sqrt(h))


def as_number(v):
    """REPD numeric fields arrive as text and carry junk. Return None on junk."""
    if v is None:
        return None
    try:
        n = float(str(v).strip())
    except ValueError:
        return None
    return n if math.isfinite(n) else None


def grid_precision_m(easting, northing):
    """The coarsest round number both coordinates sit on.

    This is a FLOOR on the uncertainty, not the uncertainty. A reference given
    to the nearest kilometre cannot be better than a kilometre; one given to
    the metre can still be a kilometre out, and at the two sites with ground
    truth it is.
    """
    for step in (1000, 500, 100, 10):
        if easting % step == 0 and northing % step == 0:
            return step
    return 1


def read_windel(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit('openpyxl is needed to read the WINDEL workbook: pip install openpyxl')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    idx = {name: k for k, name in enumerate(next(it))}
    out = {}
    for row in it:
        if not row or row[0] is None:
            continue
        ref = str(row[idx['REFID']])
        easting, northing = row[idx['EASTING']], row[idx['NORTHING']]
        # TURBHT is TIP height, inferred from how it tracks per-turbine
        # capacity: 2 MW records sit at a 115 m median and 6 to 7 MW at 200 m,
        # which are tip figures for those machines, not hub figures. The field
        # is documented only as "height of the wind turbines".
        tip = as_number(row[idx['TURBHT']])
        # A handful of records carry 1 m, 15 m and the like. No wind turbine in
        # the REPD is 1 m tall, so those are data entry errors, not small
        # turbines, and a height that cannot be true is worse than no height.
        if tip is not None and not (20 <= tip <= 400):
            tip = None
        count = as_number(row[idx['TURBNUM']])
        if count is not None and not (1 <= count <= 1000):
            count = None
        out[ref] = {
            'ref': ref,
            'name': (row[idx['SITENAME']] or '').strip(),
            'lat': row[idx['LAT']], 'lon': row[idx['LNG']],
            'mw': as_number(row[idx['INSTCAP']]) or 0.0,
            'status': row[idx['DEVSTATSHT']],
            'offshore': row[idx['TECHTYPE']] == 'Wind Offshore',
            'authority': row[idx['PLANAUTH']] or '',
            'operator': (row[idx['OPERATOR']] or '').strip(),
            'turbines': int(count) if count else None,
            'tipHeightM': round(tip) if tip else None,
            'turbineMw': as_number(row[idx['TURBCAP']]),
            'gridPrecisionM': grid_precision_m(int(easting), int(northing)),
        }
    return out


def main():
    if not os.path.exists(XLSX):
        sys.exit(f'{XLSX} not found')
    windel = read_windel(XLSX)
    mirror = {f['ref']: f for f in json.load(open(OUT))}

    merged = {}
    disagree = 0
    for ref, w in windel.items():
        merged[ref] = dict(w)
    for ref, m in mirror.items():
        w = windel.get(ref)
        if w is None:
            merged[ref] = dict(m, operator='', turbines=None, tipHeightM=None,
                               turbineMw=None, gridPrecisionM=None)
            continue
        # The mirror is the later snapshot, so its position and status win; the
        # turbine attributes only exist in WINDEL.
        d = great_circle_m(m['lat'], m['lon'], w['lat'], w['lon'])
        rec = dict(w)
        rec.update(name=m['name'], lat=m['lat'], lon=m['lon'], mw=m['mw'],
                   status=m['status'], offshore=m['offshore'],
                   authority=m['authority'] or w['authority'])
        if d > 100:
            rec['sourceDisagreementM'] = round(d)
            disagree += 1
        merged[ref] = rec

    farms = sorted(merged.values(), key=lambda f: (f['name'] or '', f['ref']))
    json.dump(farms, open(OUT, 'w'), indent=0)

    live = [f for f in farms if f['status'] in LIVE]
    with_n = [f for f in live if f.get('turbines')]
    with_h = [f for f in live if f.get('tipHeightM')]
    print(f'{OUT}: {len(farms)} records ({len(live)} live)')
    print(f'  turbine count on {len(with_n)}/{len(live)} live '
          f'({100 * len(with_n) / len(live):.0f}%), {sum(f["turbines"] for f in with_n)} turbines')
    print(f'  tip height on   {len(with_h)}/{len(live)} live '
          f'({100 * len(with_h) / len(live):.0f}%)')
    print(f'  positions disagreeing by over 100 m between the two sources: {disagree}')


if __name__ == '__main__':
    main()
